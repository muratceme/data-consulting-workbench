"""Files & self-service data: application-specific data preparation from a recipe (code) and an evidence-bound self-service
preparation plan that produces the same recipe plus the application's native artifact (agent).

Recipe vocabulary (JSON list of steps): yeniden_adlandir, kirp, buyuk_kucuk, tip, bos_doldur, filtre, tekillestir, kolon_sil, bol, birlestir.
Each product applies it with its own engine or format: Excel workbook with a table (openpyxl), Power Query M, a new Google Sheets
tab, DuckDB SQL, OpenRefine operation history, Parquet with an explicit Arrow schema. The source file is never changed."""
import csv
import datetime as dt
import io
import json
import re
import uuid
from pathlib import Path
from urllib.parse import quote, urlencode

import native
from products.core import (NativeError, agent_task, bearer, driver, fold, journal, name, names, need_list, need_text, number, opt_int, plan_json,
                           profile, read_table, receipt, required, rows_md, secret, write_artifact, write_csv, write_text)

OPS = {'yeniden_adlandir', 'kirp', 'buyuk_kucuk', 'tip', 'bos_doldur', 'filtre', 'tekillestir', 'kolon_sil', 'bol', 'birlestir'}
TYPES = {'integer', 'decimal', 'date', 'text'}
CONDITIONS = {'bos_degil', 'esit', 'esit_degil', 'buyuk', 'kucuk', 'icerir'}
COL = re.compile(r'[^\x00-\x1f"\'`\[\]]{1,128}')


# ---------- Recipe validation and local engine ----------
def check_recipe(steps, columns):
    """Validate step by step, tracking renames/new/removed columns. Raises ValueError (agent) — callers convert to NativeError."""
    if not isinstance(steps, list) or not 1 <= len(steps) <= 100:
        raise ValueError('Reçete 1–100 adımlık liste olmalı.')
    cols = list(columns)
    def need(c, i):
        if c not in cols:
            raise ValueError(f'adım {i}: o noktada olmayan kolon: {c}')
    for i, s in enumerate(steps, 1):
        if not isinstance(s, dict) or s.get('islem') not in OPS:
            raise ValueError(f'adım {i}: islem şunlardan biri olmalı: ' + ', '.join(sorted(OPS)))
        op = s['islem']
        for key in ('kolon', 'yeni'):
            if isinstance(s.get(key), str) and not COL.fullmatch(s[key]):
                raise ValueError(f'adım {i}: {key} adı geçersiz karakter içeriyor.')
        if op == 'yeniden_adlandir':
            need(s.get('kolon'), i)
            if not s.get('yeni') or s['yeni'] in cols:
                raise ValueError(f'adım {i}: yeni ad boş veya zaten var.')
            cols[cols.index(s['kolon'])] = s['yeni']
        elif op in ('kirp', 'kolon_sil', 'tekillestir', 'birlestir'):
            targets = s.get('kolonlar')
            if not isinstance(targets, list) or not targets:
                raise ValueError(f'adım {i}: kolonlar listesi gerekli.')
            for c in targets:
                need(c, i)
            if op == 'kolon_sil':
                cols = [c for c in cols if c not in targets]
            if op == 'birlestir':
                if not s.get('yeni') or s['yeni'] in cols:
                    raise ValueError(f'adım {i}: birleştirme için yeni kolon adı gerekli.')
                cols.append(s['yeni'])
        elif op == 'buyuk_kucuk':
            need(s.get('kolon'), i)
            if s.get('bicim') not in ('upper', 'lower', 'title'):
                raise ValueError(f'adım {i}: bicim upper, lower veya title olmalı.')
        elif op == 'tip':
            need(s.get('kolon'), i)
            if s.get('tip') not in TYPES:
                raise ValueError(f'adım {i}: tip integer, decimal, date veya text olmalı.')
            if s['tip'] == 'date' and s.get('bicim') and not re.fullmatch(r'[%dmyYHMS./ :-]{2,20}', s['bicim']):
                raise ValueError(f'adım {i}: tarih bicimi strptime biçimi olmalı (ör. %d.%m.%Y).')
        elif op == 'bos_doldur':
            need(s.get('kolon'), i)
            if 'deger' not in s:
                raise ValueError(f'adım {i}: deger gerekli.')
        elif op == 'filtre':
            need(s.get('kolon'), i)
            if s.get('kosul') not in CONDITIONS:
                raise ValueError(f'adım {i}: kosul şunlardan biri olmalı: ' + ', '.join(sorted(CONDITIONS)))
        elif op == 'bol':
            need(s.get('kolon'), i)
            new = s.get('yeni')
            if not isinstance(new, list) or len(new) < 2 or any(n in cols for n in new) or not s.get('ayirici'):
                raise ValueError(f'adım {i}: bol için ayirici ve en az iki yeni kolon adı gerekli.')
            cols += new
    return cols


def convert(value, kind, fmt=None):
    if value in (None, ''):
        return None
    if kind == 'text':
        return str(value)
    if kind in ('integer', 'decimal'):
        n = number(value)
        if n is None:
            raise ValueError
        return int(n) if kind == 'integer' else n
    text = str(value).strip()
    for f in ([fmt] if fmt else []) + ['%Y-%m-%d', '%d.%m.%Y', '%d/%m/%Y', '%Y/%m/%d']:
        try:
            return dt.datetime.strptime(text, f).date().isoformat()
        except ValueError:
            continue
    raise ValueError


def apply_recipe(columns, rows, steps):
    try:
        final = check_recipe(steps, columns)
    except ValueError as exc:
        raise NativeError(str(exc)) from None
    cols, data, log = list(columns), [dict(r) for r in rows], []
    for i, s in enumerate(steps, 1):
        op, before, issues = s['islem'], len(data), 0
        if op == 'yeniden_adlandir':
            for r in data:
                r[s['yeni']] = r.pop(s['kolon'], None)
            cols[cols.index(s['kolon'])] = s['yeni']
        elif op == 'kirp':
            for r in data:
                for c in s['kolonlar']:
                    if isinstance(r.get(c), str):
                        r[c] = re.sub(r'\s+', ' ', r[c]).strip()
        elif op == 'buyuk_kucuk':
            fn = {'upper': lambda t: t.replace('i', 'İ').upper(), 'lower': lambda t: t.replace('I', 'ı').replace('İ', 'i').lower(), 'title': str.title}[s['bicim']]
            for r in data:
                if isinstance(r.get(s['kolon']), str):
                    r[s['kolon']] = fn(r[s['kolon']])
        elif op == 'tip':
            for r in data:
                try:
                    r[s['kolon']] = convert(r.get(s['kolon']), s['tip'], s.get('bicim'))
                except ValueError:
                    r[s['kolon']] = None
                    issues += 1
        elif op == 'bos_doldur':
            for r in data:
                if r.get(s['kolon']) in (None, ''):
                    r[s['kolon']] = s['deger']
        elif op == 'filtre':
            c, k, v = s['kolon'], s['kosul'], s.get('deger')
            def keep(r):
                x = r.get(c)
                if k == 'bos_degil':
                    return x not in (None, '')
                if k == 'icerir':
                    return str(v).lower() in str(x or '').lower()
                if k in ('buyuk', 'kucuk'):
                    a, b = number(x), number(v)
                    return a is not None and b is not None and (a > b if k == 'buyuk' else a < b)
                return (fold(x) == fold(v)) == (k == 'esit')
            data = [r for r in data if keep(r)]
        elif op == 'tekillestir':
            seen, kept = set(), []
            for r in data:
                key = tuple(fold(r.get(c)) for c in s['kolonlar'])
                if key not in seen:
                    seen.add(key)
                    kept.append(r)
            data = kept
        elif op == 'kolon_sil':
            cols = [c for c in cols if c not in s['kolonlar']]
            for r in data:
                for c in s['kolonlar']:
                    r.pop(c, None)
        elif op == 'bol':
            for r in data:
                parts = str(r.get(s['kolon']) or '').split(s['ayirici'], len(s['yeni']) - 1)
                for n, p in zip(s['yeni'], parts + [''] * len(s['yeni'])):
                    r[n] = p
            cols += s['yeni']
        elif op == 'birlestir':
            for r in data:
                r[s['yeni']] = s.get('ayirici', ' ').join(str(r.get(c) or '') for c in s['kolonlar']).strip()
            cols.append(s['yeni'])
        log.append({'adim': i, 'islem': op, 'satir_once': before, 'satir_sonra': len(data), 'donusmeyen_deger': issues})
    assert cols == final
    return cols, [{c: r.get(c) for c in cols} for r in data], log


def load_recipe(cfg):
    path = Path(cfg.get('recipe_file', 'recete.json'))
    try:
        steps = json.loads(path.read_text(encoding='utf-8'))
    except ValueError:
        raise NativeError('recete.json geçerli JSON değil.') from None
    return steps.get('adimlar', steps) if isinstance(steps, dict) else steps


def recipe_types(steps):
    return {s['kolon']: s['tip'] for s in steps if s['islem'] == 'tip'}


# ---------- Power Query M ----------
def m_text(value):
    return '"' + str(value).replace('"', '""') + '"'


def m_query(steps, source_path, delimiter=','):
    lines = [f'    Kaynak = Csv.Document(File.Contents({m_text(source_path)}), [Delimiter={m_text(delimiter)}, Encoding=65001, QuoteStyle=QuoteStyle.Csv]),',
             '    Basliklar = Table.PromoteHeaders(Kaynak, [PromoteAllScalars=true]),']
    prev = 'Basliklar'
    for i, s in enumerate(steps, 1):
        step, op = f'Adim{i}', s['islem']
        if op == 'yeniden_adlandir':
            expr = f'Table.RenameColumns({prev}, {{{{{m_text(s["kolon"])}, {m_text(s["yeni"])}}}}})'
        elif op == 'kirp':
            expr = f'Table.TransformColumns({prev}, {{' + ', '.join(f'{{{m_text(c)}, Text.Trim, type text}}' for c in s['kolonlar']) + '})'
        elif op == 'buyuk_kucuk':
            fn = {'upper': 'Text.Upper', 'lower': 'Text.Lower', 'title': 'Text.Proper'}[s['bicim']]
            expr = f'Table.TransformColumns({prev}, {{{{{m_text(s["kolon"])}, each {fn}(_, "tr-TR"), type text}}}})'
        elif op == 'tip':
            mtype = {'integer': 'Int64.Type', 'decimal': 'type number', 'date': 'type date', 'text': 'type text'}[s['tip']]
            expr = f'Table.TransformColumnTypes({prev}, {{{{{m_text(s["kolon"])}, {mtype}}}}}, "tr-TR")'
        elif op == 'bos_doldur':
            expr = f'Table.ReplaceValue({prev}, null, {m_text(s["deger"])}, Replacer.ReplaceValue, {{{m_text(s["kolon"])}}})'
        elif op == 'filtre':
            c, v = f'[{s["kolon"]}]', s.get('deger')
            cond = {'bos_degil': f'{c} <> null and {c} <> ""', 'esit': f'{c} = {m_text(v)}', 'esit_degil': f'{c} <> {m_text(v)}', 'buyuk': f'{c} > {number(v)}',
                    'kucuk': f'{c} < {number(v)}', 'icerir': f'Text.Contains(Text.From({c}), {m_text(v)}, Comparer.OrdinalIgnoreCase)'}[s['kosul']]
            expr = f'Table.SelectRows({prev}, each {cond})'
        elif op == 'tekillestir':
            expr = f'Table.Distinct({prev}, {{' + ', '.join(m_text(c) for c in s['kolonlar']) + '})'
        elif op == 'kolon_sil':
            expr = f'Table.RemoveColumns({prev}, {{' + ', '.join(m_text(c) for c in s['kolonlar']) + '})'
        elif op == 'bol':
            expr = (f'Table.SplitColumn({prev}, {m_text(s["kolon"])}, Splitter.SplitTextByDelimiter({m_text(s["ayirici"])}, QuoteStyle.None), '
                    '{' + ', '.join(m_text(n) for n in s['yeni']) + '})')
        else:
            expr = (f'Table.AddColumn({prev}, {m_text(s["yeni"])}, each Text.Combine({{' + ', '.join(f'Text.From([{c}])' for c in s['kolonlar']) +
                    f'}}, {m_text(s.get("ayirici", " "))}), type text)')
        lines.append(f'    {step} = {expr},')
        prev = step
    lines[-1] = lines[-1].rstrip(',')
    return 'let\n' + '\n'.join(lines) + f'\nin\n    {prev}\n'


# ---------- DuckDB SQL ----------
def q(c):
    return '"' + c.replace('"', '""') + '"'


def duckdb_sql(steps, source, columns):
    cols, ctes, prev = list(columns), [], 'kaynak'
    ctes.append(f"kaynak AS (SELECT * FROM read_csv_auto('{source.replace(chr(39), chr(39) * 2)}', header = true, all_varchar = true))")
    for i, s in enumerate(steps, 1):
        op, name_ = s['islem'], f'adim{i}'
        select = None
        if op == 'yeniden_adlandir':
            select = ', '.join(f'{q(c)} AS {q(s["yeni"])}' if c == s['kolon'] else q(c) for c in cols)
            cols[cols.index(s['kolon'])] = s['yeni']
        elif op == 'kirp':
            select = ', '.join(f"trim(regexp_replace({q(c)}, '\\s+', ' ', 'g')) AS {q(c)}" if c in s['kolonlar'] else q(c) for c in cols)
        elif op == 'buyuk_kucuk':
            fn = {'upper': 'upper({c})', 'lower': 'lower({c})',
                  'title': "array_to_string(list_transform(string_split(lower({c}), ' '), lambda w: upper(w[1]) || w[2:]), ' ')"}[s['bicim']]
            select = ', '.join(fn.format(c=q(c)) + f' AS {q(c)}' if c == s['kolon'] else q(c) for c in cols)
        elif op == 'tip':
            expr = {'integer': "TRY_CAST(replace({c}, '.', '') AS BIGINT)", 'decimal': "TRY_CAST(replace(replace({c}, '.', ''), ',', '.') AS DOUBLE)",
                    'date': ("CAST(try_strptime({c}, '" + (s.get('bicim') or '%d.%m.%Y') + "') AS DATE)"), 'text': 'CAST({c} AS VARCHAR)'}[s['tip']]
            select = ', '.join(expr.format(c=q(c)) + f' AS {q(c)}' if c == s['kolon'] else q(c) for c in cols)
        elif op == 'bos_doldur':
            select = ', '.join(f"coalesce(nullif({q(c)}, ''), '{str(s['deger']).replace(chr(39), chr(39) * 2)}') AS {q(c)}" if c == s['kolon'] else q(c) for c in cols)
        elif op == 'filtre':
            c, v = q(s['kolon']), str(s.get('deger', '')).replace("'", "''")
            cond = {'bos_degil': f"{c} IS NOT NULL AND {c} <> ''", 'esit': f"{c} = '{v}'", 'esit_degil': f"{c} <> '{v}'", 'buyuk': f'TRY_CAST({c} AS DOUBLE) > {number(v)}',
                    'kucuk': f'TRY_CAST({c} AS DOUBLE) < {number(v)}', 'icerir': f"{c} ILIKE '%{v}%'"}[s['kosul']]
            ctes.append(f'{name_} AS (SELECT * FROM {prev} WHERE {cond})')
            prev = name_
            continue
        elif op == 'tekillestir':
            ctes.append(f'{name_} AS (SELECT * FROM {prev} QUALIFY row_number() OVER (PARTITION BY ' + ', '.join(f'lower({q(c)})' for c in s['kolonlar']) + ') = 1)')
            prev = name_
            continue
        elif op == 'kolon_sil':
            cols = [c for c in cols if c not in s['kolonlar']]
            select = ', '.join(q(c) for c in cols)
        elif op == 'bol':
            parts = ', '.join(f"split_part({q(s['kolon'])}, '{s['ayirici'].replace(chr(39), chr(39) * 2)}', {n + 1}) AS {q(new)}" for n, new in enumerate(s['yeni']))
            select = ', '.join(q(c) for c in cols) + ', ' + parts
            cols += s['yeni']
        else:
            select = ', '.join(q(c) for c in cols) + f", concat_ws('{s.get('ayirici', ' ').replace(chr(39), chr(39) * 2)}', " + ', '.join(q(c) for c in s['kolonlar']) + f') AS {q(s["yeni"])}'
            cols.append(s['yeni'])
        ctes.append(f'{name_} AS (SELECT {select} FROM {prev})')
        prev = name_
    return 'WITH ' + ',\n'.join(ctes) + f'\nSELECT * FROM {prev}'


# ---------- OpenRefine operation history ----------
def refine_ops(steps):
    engine = {'facets': [], 'mode': 'row-based'}
    ops = []
    for i, s in enumerate(steps, 1):
        op = s['islem']
        if op == 'yeniden_adlandir':
            ops.append({'op': 'core/column-rename', 'oldColumnName': s['kolon'], 'newColumnName': s['yeni'], 'description': f'Rename {s["kolon"]}'})
        elif op in ('kirp', 'buyuk_kucuk', 'tip', 'bos_doldur'):
            targets = s['kolonlar'] if op == 'kirp' else [s['kolon']]
            grel = {'kirp': 'value.trim()', 'buyuk_kucuk': {'upper': 'value.toUppercase()', 'lower': 'value.toLowercase()', 'title': 'value.toTitlecase()'}.get(s.get('bicim')),
                    'tip': {'integer': 'value.toNumber()', 'decimal': 'value.replace(".", "").replace(",", ".").toNumber()', 'date': 'value.toDate()', 'text': 'value.toString()'}.get(s.get('tip')),
                    'bos_doldur': 'if(isBlank(value), ' + json.dumps(str(s.get('deger'))) + ', value)'}[op]
            for c in targets:
                ops.append({'op': 'core/text-transform', 'engineConfig': engine, 'columnName': c, 'expression': 'grel:' + grel, 'onError': 'keep-original', 'repeat': False, 'repeatCount': 10,
                            'description': f'Text transform on {c}'})
        elif op == 'kolon_sil':
            ops += [{'op': 'core/column-removal', 'columnName': c, 'description': f'Remove {c}'} for c in s['kolonlar']]
        elif op == 'bol':
            ops.append({'op': 'core/column-split', 'engineConfig': engine, 'columnName': s['kolon'], 'guessCellType': False, 'removeOriginalColumn': False, 'mode': 'separator',
                        'separator': s['ayirici'], 'regex': False, 'maxColumns': len(s['yeni']), 'description': f'Split {s["kolon"]}'})
        elif op == 'birlestir':
            expr = ' + ' + json.dumps(s.get('ayirici', ' ')) + ' + '
            ops.append({'op': 'core/column-addition', 'engineConfig': engine, 'baseColumnName': s['kolonlar'][0], 'newColumnName': s['yeni'], 'columnInsertIndex': 0, 'onError': 'set-to-blank',
                        'expression': 'grel:' + expr.join(f'cells[{json.dumps(c)}].value' for c in s['kolonlar']), 'description': f'Combine into {s["yeni"]}'})
        else:
            raise NativeError(f'adım {i}: {op} OpenRefine işlem geçmişinde facet gerektirir; bu pakette desteklenmez (DuckDB/Excel paketlerini kullanın).')
    return ops


# ---------- Readers per product ----------
def read_excel(cfg):
    openpyxl = driver('openpyxl')
    wb = openpyxl.load_workbook(required(cfg, 'input_file'), read_only=True, data_only=True)
    ws = wb[cfg['sheet']] if cfg.get('sheet') else wb.worksheets[0]
    values = list(ws.iter_rows(values_only=True))
    wb.close()
    header_row = opt_int(cfg, 'header_row', 1, 1, 50)
    header = [str(h).strip() if h is not None else f'kolon_{i + 1}' for i, h in enumerate(values[header_row - 1])]
    rows = [{h: (v.isoformat() if isinstance(v, (dt.date, dt.datetime)) else v) for h, v in zip(header, r)} for r in values[header_row:] if any(x not in (None, '') for x in r)]
    return header, rows


def sheets_client(factory):
    return bearer(factory, 'https://sheets.googleapis.com', 'GOOGLE_ACCESS_TOKEN')


def read_sheet(cfg, factory):
    client = sheets_client(factory)
    sid = name(cfg, 'spreadsheet_id', re.compile(r'[A-Za-z0-9_-]{20,100}'))
    values = client.json('GET', f'/v4/spreadsheets/{sid}/values/{quote(required(cfg, "range"), safe="")}?valueRenderOption=FORMATTED_VALUE').get('values', [])
    if not values:
        raise NativeError('Aralık boş.')
    header = [str(h).strip() for h in values[0]]
    return header, [dict(zip(header, r + [''] * (len(header) - len(r)))) for r in values[1:]]


def read_any(key, cfg, factory):
    if key == 'excel':
        return read_excel(cfg)
    if key == 'sheets':
        return read_sheet(cfg, factory)
    return read_table(required(cfg, 'input_file'))


# ---------- Code task per product ----------
def write_workbook(out, columns, rows, types):
    openpyxl = driver('openpyxl')
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Hazir'
    ws.append(columns)
    for r in rows:
        ws.append([(dt.date.fromisoformat(r[c]) if types.get(c) == 'date' and r.get(c) else r.get(c)) for c in columns])
    if rows:
        ref = f'A1:{get_column_letter(len(columns))}{len(rows) + 1}'
        table = Table(displayName='HazirVeri', ref=ref)
        table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(table)
    for i, c in enumerate(columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(str(c)) + 2))
        if types.get(c) == 'date':
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for x in cell:
                    x.number_format = 'DD.MM.YYYY'
    buf = io.BytesIO()
    wb.save(buf)
    return write_artifact(out, 'hazir_veri.xlsx', buf.getvalue())


def prepare(cfg, out, factory):
    key = cfg['operation'].rsplit('-data-prep', 1)[0]
    label = PRODUCTS[key]
    steps = load_recipe(cfg)
    if key == 'openrefine':
        return openrefine_prepare(cfg, out, factory, steps)
    columns, rows = read_any(key, cfg, factory)
    if key == 'duckdb':
        duckdb = driver('duckdb')
        try:
            check_recipe(steps, columns)
        except ValueError as exc:
            raise NativeError(str(exc)) from None
        sql = duckdb_sql(steps, str(Path(cfg['input_file']).resolve()), columns)
        write_text(out, 'hazirlama.sql', sql)
        target = out / ('hazir_veri.' + cfg.get('output_format', 'parquet'))
        con = duckdb.connect()
        try:
            fmt = 'PARQUET' if cfg.get('output_format', 'parquet') == 'parquet' else 'CSV, HEADER'
            con.execute(f"COPY ({sql}) TO '{str(target).replace(chr(39), chr(39) * 2)}' (FORMAT {fmt})")
            count = con.execute(f'SELECT COUNT(*) FROM ({sql})').fetchone()[0]
        except duckdb.Error as exc:
            raise NativeError('DuckDB sorgusu çalışmadı: ' + str(exc).splitlines()[0][:200]) from None
        finally:
            con.close()
        journal(out, state='verified')
        return receipt(out, 'hazirlama.json', {'urun': label, 'kaynak_satir': len(rows), 'cikti_satir': count, 'cikti': target.name, 'motor': 'DuckDB'})
    final_cols, data, log = apply_recipe(columns, rows, steps)
    write_csv(out, 'adim_gunlugu.csv', log, ['adim', 'islem', 'satir_once', 'satir_sonra', 'donusmeyen_deger'])
    types = recipe_types(steps)
    if key == 'excel':
        artifact = write_workbook(out, final_cols, data, types)
    elif key == 'power-query':
        write_text(out, 'sorgu.pq', m_query(steps, str(Path(cfg['input_file']).resolve()), cfg.get('delimiter', ',')))
        artifact = write_csv(out, 'onizleme.csv', data, final_cols)
    elif key == 'sheets':
        artifact = sheets_write(cfg, factory, out, final_cols, data)
    else:
        artifact = parquet_write(out, final_cols, data, types, cfg)
    journal(out, state='verified')
    return receipt(out, 'hazirlama.json', {'urun': label, 'kaynak_satir': len(rows), 'cikti_satir': len(data), 'kolonlar': final_cols, 'cikti': artifact,
                                           'donusmeyen_deger': sum(l['donusmeyen_deger'] for l in log), 'not': 'Kaynak değiştirilmedi.'})


def sheets_write(cfg, factory, out, columns, rows):
    client = sheets_client(factory)
    sid = cfg['spreadsheet_id']
    title = cfg.get('output_sheet') or 'datawork_' + dt.datetime.now().strftime('%Y%m%d_%H%M')
    if not re.fullmatch(r"[^'\[\]\\/?*:]{1,90}", title):
        raise NativeError('output_sheet geçersiz sekme adı.')
    meta = client.json('GET', f'/v4/spreadsheets/{sid}?fields=sheets.properties.title')
    if title in [s.get('properties', {}).get('title') for s in meta.get('sheets', [])]:
        raise NativeError('Bu adla sekme zaten var; mevcut sekme üzerine yazılmaz.')
    client.json('POST', f'/v4/spreadsheets/{sid}:batchUpdate', {'requests': [{'addSheet': {'properties': {'title': title}}}]})
    journal(out, created_sheet=title)
    values = [columns] + [[('' if r.get(c) is None else r.get(c)) for c in columns] for r in rows]
    client.json('PUT', f"/v4/spreadsheets/{sid}/values/{quote(title + '!A1', safe='')}?valueInputOption=RAW", {'range': f'{title}!A1', 'majorDimension': 'ROWS', 'values': values})
    back = client.json('GET', f"/v4/spreadsheets/{sid}/values/{quote(title + '!A:A', safe='')}").get('values', [])
    if len(back) != len(values):
        raise NativeError('Yazılan satır sayısı geri okumada tutmadı; sekme adı run-state.json içinde.')
    return {'sekme': title, 'satir': len(rows)}


def parquet_write(out, columns, rows, types, cfg):
    pa = driver('pyarrow')
    pq = driver('pyarrow.parquet')
    arrow = {'integer': pa.int64(), 'decimal': pa.float64(), 'date': pa.date32(), 'text': pa.string()}
    fields = [pa.field(c, arrow.get(types.get(c), pa.string())) for c in columns]
    arrays = []
    for f in fields:
        vals = [r.get(f.name) for r in rows]
        if f.type == pa.date32():
            vals = [dt.date.fromisoformat(v) if v else None for v in vals]
        elif f.type == pa.string():
            vals = [None if v is None else str(v) for v in vals]
        arrays.append(pa.array(vals, type=f.type))
    table = pa.Table.from_arrays(arrays, schema=pa.schema(fields))
    buf = io.BytesIO()
    pq.write_table(table, buf, compression=cfg.get('compression', 'zstd'))
    raw = buf.getvalue()
    meta = pq.ParquetFile(io.BytesIO(raw)).metadata
    if meta.num_rows != len(rows):
        raise NativeError('Parquet geri okumada satır sayısı tutmadı.')
    artifact = write_artifact(out, 'hazir_veri.parquet', raw)
    receipt(out, 'parquet_semasi.json', {'kolonlar': [{'ad': f.name, 'tip': str(f.type)} for f in fields], 'satir': meta.num_rows, 'sikistirma': cfg.get('compression', 'zstd')})
    return artifact


def openrefine_prepare(cfg, out, factory, steps):
    columns, rows = read_table(required(cfg, 'input_file'))
    try:
        check_recipe(steps, columns)
    except ValueError as exc:
        raise NativeError(str(exc)) from None
    ops = refine_ops(steps)
    write_text(out, 'openrefine_islemleri.json', json.dumps(ops, ensure_ascii=False, indent=2))
    client = factory(cfg.get('server', 'http://127.0.0.1:3333'), {})
    token = client.json('GET', '/command/core/get-csrf-token').get('token')
    data = Path(cfg['input_file']).read_bytes()
    boundary = 'datawork' + uuid.uuid4().hex
    parts = [f'--{boundary}\r\nContent-Disposition: form-data; name="project-file"; filename="veri.csv"\r\nContent-Type: text/csv\r\n\r\n'.encode() + data + b'\r\n']
    for k, v in (('project-name', f'datawork {dt.date.today()}'), ('format', 'text/line-based/*sv'), ('options', json.dumps({'separator': ',', 'headerLines': 1}))):
        parts.append(f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode('utf-8'))
    raw, headers, status = client.request('POST', f'/command/core/create-project-from-upload?csrf_token={quote(token or "")}', raw=b''.join(parts) + f'--{boundary}--\r\n'.encode(),
                                          content_type='multipart/form-data; boundary=' + boundary, with_status=True, allow_status=(302,))
    location = native.header(headers, 'Location')
    project = re.search(r'project=(\d+)', location or '')
    if not project:
        raise NativeError('OpenRefine proje kimliği dönmedi.')
    pid = project.group(1)
    journal(out, openrefine_project=pid)
    token = client.json('GET', '/command/core/get-csrf-token').get('token')
    form = urlencode({'project': pid, 'operations': json.dumps(ops), 'csrf_token': token}).encode()
    applied = native.json_object(client.request('POST', '/command/core/apply-operations', raw=form, content_type='application/x-www-form-urlencoded')[0])
    if applied.get('code') not in ('ok', 'pending'):
        raise NativeError('OpenRefine işlemleri uygulayamadı.')
    raw, _ = client.request('POST', '/command/core/export-rows', raw=urlencode({'project': pid, 'format': 'csv'}).encode(), content_type='application/x-www-form-urlencoded',
                            accept='application/octet-stream')
    artifact = write_artifact(out, 'hazir_veri.csv', raw)
    journal(out, state='verified')
    return receipt(out, 'hazirlama.json', {'urun': 'OpenRefine', 'proje': pid, 'islem': len(ops), 'kaynak_satir': len(rows), 'cikti': artifact,
                                           'not': 'Proje OpenRefine’da kalır (geçmişi incelenebilir); kaynak dosya değiştirilmedi.'})


PRODUCTS = {'excel': 'Microsoft Excel', 'power-query': 'Power Query', 'sheets': 'Google Sheets', 'duckdb': 'DuckDB', 'openrefine': 'OpenRefine', 'arrow-parquet': 'Apache Arrow / Parquet'}


# ---------- Self-service preparation plan agent ----------
SCHEMA = ('{"sorunlar":[{"kolon":"kanıttaki kolon","sorun":"boşluk/tutarsız tarih/karışık büyük-küçük harf/tekrar","kanit":"profil bulgusu"}],'
          '"recete":[{"islem":"yeniden_adlandir|kirp|buyuk_kucuk|tip|bos_doldur|filtre|tekillestir|kolon_sil|bol|birlestir","kolon":"...","kolonlar":["..."],"yeni":"...",'
          '"tip":"integer|decimal|date|text","bicim":"...","kosul":"bos_degil|esit|esit_degil|buyuk|kucuk|icerir","deger":"...","ayirici":"..."}],"sorular":["..."]}')


def text_issues(columns, rows):
    found = []
    for c in columns:
        vals = [r.get(c) for r in rows[:5000] if isinstance(r.get(c), str) and r.get(c)]
        if not vals:
            continue
        spaces = sum(1 for v in vals if v != v.strip() or '  ' in v)
        cases = len({v.strip() for v in vals}) - len({fold(v) for v in vals})
        dates = {('nokta' if re.fullmatch(r'\d{1,2}\.\d{1,2}\.\d{4}', v) else 'tire' if re.fullmatch(r'\d{4}-\d{2}-\d{2}', v) else 'egik' if re.fullmatch(r'\d{1,2}/\d{1,2}/\d{4}', v) else None)
                 for v in vals} - {None}
        decimal_comma = sum(1 for v in vals if re.fullmatch(r'-?\d{1,3}(\.\d{3})*,\d+', v))
        found.append({'kolon': c, 'bosluk_sorunlu': spaces, 'buyuk_kucuk_varyanti': max(0, cases), 'tarih_bicimleri': sorted(dates), 'virgullu_ondalik': decimal_comma})
    return found


def prep_plan(cfg, out, factory):
    key = cfg['operation'].rsplit('-prep-plan', 1)[0]
    label = PRODUCTS[key]
    columns, rows = read_any(key, cfg, factory)
    evidence = {'kolonlar': columns, 'satir': len(rows), 'profil': profile(columns, rows), 'metin_bulgulari': text_issues(columns, rows),
                'tekrar_satir': len(rows) - len({json.dumps(r, sort_keys=True, default=str) for r in rows})}

    def validate(obj):
        for i, s in enumerate(obj.get('sorunlar') or [], 1):
            if not isinstance(s, dict) or s.get('kolon') not in columns:
                raise ValueError(f'sorun {i}: kanıtta olmayan kolon.')
        steps = obj.get('recete')
        check_recipe(steps, columns)
        if key == 'openrefine':
            try:
                refine_ops(steps)
            except NativeError as exc:
                raise ValueError(str(exc)) from None
        return obj

    def render(result):
        steps = result['recete']
        files = {'recete.json': json.dumps(steps, ensure_ascii=False, indent=2),
                 'hazirlama_plani.md': rows_md('Self-service veri hazırlama planı — ' + label,
                                               [('Sorunlar', result.get('sorunlar') or [], ['kolon', 'sorun', 'kanit']),
                                                ('Reçete', [{'adim': i + 1, **{k: (', '.join(v) if isinstance(v, list) else v) for k, v in s.items()}} for i, s in enumerate(steps)],
                                                 ['adim', 'islem', 'kolon', 'kolonlar', 'yeni', 'tip', 'bicim', 'kosul', 'deger', 'ayirici']),
                                                ('Sorular', [{'soru': q_} for q_ in result.get('sorular', [])], ['soru'])])}
        source = str(cfg.get('input_file') or 'veri.csv')
        if key in ('power-query', 'excel'):
            files['sorgu.pq'] = m_query(steps, source, cfg.get('delimiter', ','))
        if key == 'duckdb':
            files['hazirlama.sql'] = duckdb_sql(steps, source, columns) + ';\n'
        if key == 'openrefine':
            files['openrefine_islemleri.json'] = json.dumps(refine_ops(steps), ensure_ascii=False, indent=2)
        _, preview, _ = apply_recipe(columns, rows[:200], steps)
        buf = io.StringIO()
        if preview:
            w = csv.DictWriter(buf, fieldnames=list(preview[0]), lineterminator='\n')
            w.writeheader(); w.writerows(preview)
        files['onizleme_ilk200.csv'] = buf.getvalue() or 'bilgi\nboş\n'
        return files
    return agent_task(cfg, out, product=label, title='Self-service veri hazırlama planı', evidence=evidence, schema=SCHEMA, validate=validate, render=render,
                      rules='Reçeteyi yalnızca izinli işlemlerle ve sırayla kur; yeniden adlandırılan kolonu sonraki adımlarda yeni adıyla kullan. '
                            'Türkçe sayı ve tarih biçimlerini (1.234,56 ve gg.aa.yyyy) tip adımlarında belirt. Satır silen adımları gerekçelendir.')


HANDLERS = {}
for _key in PRODUCTS:
    HANDLERS[_key + '-data-prep'] = prepare
    HANDLERS[_key + '-prep-plan'] = prep_plan
