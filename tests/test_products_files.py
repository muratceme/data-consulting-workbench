"""Files & self-service product packages: one recipe applied by each application's engine/format, and the preparation plan agent."""
import csv
import importlib.util
import json
import sys
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_products import Base, model_reply  # noqa: E402
import native  # noqa: E402
from products import files  # noqa: E402

DATA = 'siparis_no,musteri,tutar,tarih\n1,  ayşe   yılmaz ,"1.234,50",01.10.2026\n1,Ayşe Yılmaz,"1.234,50",01.10.2026\n2,MEHMET,abc,32.13.2026\n'
RECIPE = [{'islem': 'kirp', 'kolonlar': ['musteri']}, {'islem': 'buyuk_kucuk', 'kolon': 'musteri', 'bicim': 'title'}, {'islem': 'tip', 'kolon': 'tutar', 'tip': 'decimal'},
          {'islem': 'tip', 'kolon': 'tarih', 'tip': 'date', 'bicim': '%d.%m.%Y'}, {'islem': 'tekillestir', 'kolonlar': ['siparis_no']},
          {'islem': 'yeniden_adlandir', 'kolon': 'musteri', 'yeni': 'musteri_adi'}]


class Engine(unittest.TestCase):
    def test_recipe_engine(self):
        cols, rows, log = files.apply_recipe(['siparis_no', 'musteri', 'tutar', 'tarih'], [dict(r) for r in csv.DictReader(DATA.splitlines())], RECIPE)
        self.assertEqual(cols, ['siparis_no', 'musteri_adi', 'tutar', 'tarih'])
        self.assertEqual(rows[0], {'siparis_no': '1', 'musteri_adi': 'Ayşe Yılmaz', 'tutar': 1234.5, 'tarih': '2026-10-01'})
        self.assertEqual((len(rows), log[2]['donusmeyen_deger'], log[3]['donusmeyen_deger']), (2, 1, 1))

    def test_recipe_tracks_renames(self):
        with self.assertRaises(ValueError):
            files.check_recipe([{'islem': 'yeniden_adlandir', 'kolon': 'a', 'yeni': 'b'}, {'islem': 'kirp', 'kolonlar': ['a']}], ['a'])

    def test_m_and_duckdb_translation(self):
        m = files.m_query(RECIPE, 'C:/veri.csv')
        self.assertIn('Table.TransformColumnTypes(Adim2, {{"tutar", type number}}, "tr-TR")', m)
        self.assertIn('Table.Distinct(Adim4, {"siparis_no"})', m)
        sql = files.duckdb_sql(RECIPE, 'veri.csv', ['siparis_no', 'musteri', 'tutar', 'tarih'])
        self.assertIn("try_strptime(\"tarih\", '%d.%m.%Y')", sql)
        self.assertIn('QUALIFY row_number() OVER (PARTITION BY lower("siparis_no")) = 1', sql)

    def test_openrefine_rejects_facet_steps(self):
        ops = files.refine_ops(RECIPE[:4])
        self.assertEqual(ops[0]['expression'], 'grel:value.trim()')
        with self.assertRaises(native.NativeError):
            files.refine_ops(RECIPE)


class Products(Base):
    def setUp(self):
        super().setUp()
        (self.root / 'veri.csv').write_text(DATA, encoding='utf-8')
        (self.root / 'recete.json').write_text(json.dumps(RECIPE), encoding='utf-8')

    def test_power_query_writes_m_and_preview(self):
        self.run_op('power-query-data-prep', input_file='veri.csv')
        self.assertIn('let\n', self.read('sorgu.pq'))
        self.assertEqual(self.read('hazirlama.json')['cikti_satir'], 2)
        self.assertEqual((self.root / 'veri.csv').read_text(encoding='utf-8'), DATA)

    @unittest.skipUnless(importlib.util.find_spec('openpyxl'), 'openpyxl kurulu değil')
    def test_excel_roundtrip(self):
        import openpyxl
        wb = openpyxl.Workbook(); ws = wb.active
        for r in csv.reader(DATA.splitlines()):
            ws.append(r)
        wb.save(self.root / 'veri.xlsx')
        self.run_op('excel-data-prep', input_file='veri.xlsx')
        out = openpyxl.load_workbook(self.out / 'hazir_veri.xlsx')
        self.assertIn('HazirVeri', out['Hazir'].tables)
        self.assertEqual(out['Hazir']['C2'].value, 1234.5)

    @unittest.skipUnless(importlib.util.find_spec('duckdb'), 'duckdb kurulu değil')
    def test_duckdb_real(self):
        self.run_op('duckdb-data-prep', input_file='veri.csv', output_format='csv')
        self.assertEqual(self.read('hazirlama.json')['cikti_satir'], 2)

    @unittest.skipUnless(importlib.util.find_spec('pyarrow'), 'pyarrow kurulu değil')
    def test_parquet_schema(self):
        self.run_op('arrow-parquet-data-prep', input_file='veri.csv')
        schema = {c['ad']: c['tip'] for c in self.read('parquet_semasi.json')['kolonlar']}
        self.assertEqual((schema['tutar'], schema['tarih']), ('double', 'date32[day]'))

    def test_sheets_new_tab_never_overwrites(self):
        sid = 'a' * 30
        self.server.add('GET', rf'/v4/spreadsheets/{sid}/values/Sayfa1%21A%3AZ\?.*', {'values': [r for r in csv.reader(DATA.splitlines())]})
        self.server.add('GET', rf'/v4/spreadsheets/{sid}\?fields=.*', {'sheets': [{'properties': {'title': 'Sayfa1'}}]})
        self.server.add('POST', rf'/v4/spreadsheets/{sid}:batchUpdate', {})
        self.server.add('PUT', rf'/v4/spreadsheets/{sid}/values/temiz%21A1\?valueInputOption=RAW', {})
        self.server.add('GET', rf'/v4/spreadsheets/{sid}/values/temiz%21A%3AA', {'values': [['siparis_no'], ['1'], ['2']]})
        self.run_op('sheets-data-prep', spreadsheet_id=sid, range='Sayfa1!A:Z', output_sheet='temiz')
        self.assertEqual(self.server.received[2][2], {'requests': [{'addSheet': {'properties': {'title': 'temiz'}}}]})
        self.out = self.root / 'again'
        self.server.add('GET', rf'/v4/spreadsheets/{sid}/values/Sayfa1%21A%3AZ\?.*', {'values': [r for r in csv.reader(DATA.splitlines())]})
        self.server.add('GET', rf'/v4/spreadsheets/{sid}\?fields=.*', {'sheets': [{'properties': {'title': 'temiz'}}]})
        with self.assertRaises(native.NativeError):
            self.run_op('sheets-data-prep', spreadsheet_id=sid, range='Sayfa1!A:Z', output_sheet='temiz')

    def test_openrefine_project_flow(self):
        (self.root / 'recete.json').write_text(json.dumps(RECIPE[:4]), encoding='utf-8')
        self.server.add('GET', '/command/core/get-csrf-token', {'token': 'c1'})
        self.server.add('POST', r'/command/core/create-project-from-upload\?csrf_token=c1', b'', 302, {'Location': 'http://127.0.0.1:3333/project?project=1234'})
        self.server.add('GET', '/command/core/get-csrf-token', {'token': 'c2'})
        self.server.add('POST', '/command/core/apply-operations', {'code': 'ok'})
        self.server.add('POST', '/command/core/export-rows', b'siparis_no,musteri\n1,Ayse\n')
        self.run_op('openrefine-data-prep', input_file='veri.csv', server='https://refine.example')
        self.assertIn('csrf_token=c2', self.server.received[3][2] if isinstance(self.server.received[3][2], str) else '')
        self.assertEqual(self.read('hazirlama.json')['proje'], '1234')


class Plan(Base):
    def test_prep_plan_validates_recipe_and_emits_sql(self):
        (self.root / 'veri.csv').write_text(DATA, encoding='utf-8')
        bad = {'recete': [{'islem': 'kirp', 'kolonlar': ['adres']}]}
        good = {'sorunlar': [{'kolon': 'musteri', 'sorun': 'boşluk', 'kanit': 'bosluk_sorunlu=1'}], 'recete': RECIPE}
        with model_reply(bad, good):
            self.run_op('duckdb-prep-plan', input_file='veri.csv')
        self.assertIn('WITH kaynak AS', self.read('hazirlama.sql'))
        self.assertIn('musteri_adi', self.read('onizleme_ilk200.csv'))
        self.assertEqual(json.loads((self.out / 'recete.json').read_text(encoding='utf-8')), RECIPE)
